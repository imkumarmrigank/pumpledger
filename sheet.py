"""Read/write the Bajrang Petroleum daily-ledger workbook.

The workbook has one sheet, one row per day (row 3 = day 1, row N+2 = day N),
and a fixed set of columns A..Y. Columns E, H, L, V, W, X, Y are formulas in
the original template (e.g. E = C*D, V = SUM(N:U), W = M-V, Y = L-X) and are
never touched here — only the "input" columns are written, exactly the way
the xlsx-editing convention expects: find the designated input cells, write
only there, leave every existing formula untouched.
"""
from __future__ import annotations

import shutil
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

COL = {
    'day': 'A', 'ob': 'B', 'hsd_unit': 'C', 'hsd_rate': 'D', 'hsd_amt': 'E',
    'ms_unit': 'F', 'ms_rate': 'G', 'ms_amt': 'H', 'lub': 'I', 'coffee': 'J',
    'collection': 'K', 'total_inflow': 'L', 'total_exp': 'M',
    'bank': 'N', 'ptm': 'O', 'upi': 'P', 'tsale': 'Q', 'fleet': 'R',
    'rbabu': 'S', 'ranjit': 'T', 'others': 'U',
    'v_sum': 'V', 'w_pexp': 'W', 'x_total': 'X', 'balance': 'Y',
}
FIRST_DAY_ROW = 3


def row_for_day(day: int) -> int:
    return FIRST_DAY_ROW + day - 1


@dataclass
class DayEntry:
    day: int
    opening_cash: float | None = None   # only used for day 1 (B is a formula for every other day)
    hsd_unit: float | None = None
    hsd_rate: float | None = None
    ms_unit: float | None = None
    ms_rate: float | None = None
    lub: float = 0.0
    coffee: float = 0.0
    collection: float = 0.0
    total_expenses: float | None = None   # M — the slip's stated "T-Exp" / total expenses
    bank: float = 0.0
    ptm: float = 0.0
    upi: float = 0.0
    tsale: float = 0.0
    fleet: float = 0.0
    rbabu: float = 0.0
    ranjit: float = 0.0
    others: float = 0.0
    source_images: list = field(default_factory=list)


def load_workbook(path: str):
    return openpyxl.load_workbook(path, data_only=False)


def load_existing_snapshot(path: str) -> dict[int, dict]:
    """Cached (already-computed) values per day from the ORIGINAL file, used
    as the 'what the sheet currently says' baseline for reconciliation."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    snap = {}
    for day in range(1, 32):
        row = row_for_day(day)
        vals = {name: ws[f'{letter}{row}'].value for name, letter in COL.items()}
        if all(v is None for k, v in vals.items() if k != 'day'):
            continue
        snap[day] = vals
    return snap


def write_day(ws, entry: DayEntry) -> None:
    row = row_for_day(entry.day)
    ws[f"{COL['day']}{row}"] = entry.day
    if entry.day == 1 and entry.opening_cash is not None:
        ws[f"{COL['ob']}{row}"] = entry.opening_cash
    if entry.hsd_unit is not None:
        ws[f"{COL['hsd_unit']}{row}"] = entry.hsd_unit
    if entry.hsd_rate is not None:
        ws[f"{COL['hsd_rate']}{row}"] = entry.hsd_rate
    if entry.ms_unit is not None:
        ws[f"{COL['ms_unit']}{row}"] = entry.ms_unit
    if entry.ms_rate is not None:
        ws[f"{COL['ms_rate']}{row}"] = entry.ms_rate
    ws[f"{COL['lub']}{row}"] = entry.lub or None
    ws[f"{COL['coffee']}{row}"] = entry.coffee or None
    ws[f"{COL['collection']}{row}"] = entry.collection or None
    if entry.total_expenses is not None:
        ws[f"{COL['total_exp']}{row}"] = entry.total_expenses
    ws[f"{COL['bank']}{row}"] = entry.bank or None
    ws[f"{COL['ptm']}{row}"] = entry.ptm or None
    ws[f"{COL['upi']}{row}"] = entry.upi or None
    ws[f"{COL['tsale']}{row}"] = entry.tsale or None
    ws[f"{COL['fleet']}{row}"] = entry.fleet or None
    ws[f"{COL['rbabu']}{row}"] = entry.rbabu or None
    ws[f"{COL['ranjit']}{row}"] = entry.ranjit or None
    ws[f"{COL['others']}{row}"] = entry.others or None


def build_reconciled_workbook(base_path: str, out_path: str, entries: dict[int, DayEntry]) -> None:
    """Copy the original file (preserving every formula/format we don't
    touch) and overwrite the input cells for each reconciled day."""
    shutil.copyfile(base_path, out_path)
    wb = load_workbook(out_path)
    ws = wb[wb.sheetnames[0]]
    for entry in entries.values():
        write_day(ws, entry)
    wb.save(out_path)


def computed_inflow_total(entry: DayEntry, opening_cash_for_calc: float) -> float:
    hsd_amt = (entry.hsd_unit or 0) * (entry.hsd_rate or 0)
    ms_amt = (entry.ms_unit or 0) * (entry.ms_rate or 0)
    return opening_cash_for_calc + hsd_amt + ms_amt + entry.lub + entry.coffee + entry.collection


def computed_named_total(entry: DayEntry) -> float:
    return (entry.bank + entry.ptm + entry.upi + entry.tsale + entry.fleet
            + entry.rbabu + entry.ranjit + entry.others)


def computed_balance(entry: DayEntry, opening_cash_for_calc: float) -> float:
    inflow = computed_inflow_total(entry, opening_cash_for_calc)
    named = computed_named_total(entry)
    total_exp = entry.total_expenses or 0.0
    pexp = total_exp - named
    total_out = named + pexp
    return inflow - total_out
